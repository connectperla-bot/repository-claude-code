#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Il foglio dei margini: ogni variante in vendita, con quanto resta in tasca.

PERCHE' ESISTE
perla-verifica-margini.py risponde a UNA domanda -- «c'e' qualcosa sotto il
venti per cento?» -- e la risponde bene, ma sul terminale e per venticinque
righe. La proprietaria decide i prezzi, e per decidere serve vedere tutto
insieme e sapere SU COSA si puo' agire: se un margine e' basso perche' costa
il prodotto o perche' costa la spedizione sono due decisioni diverse.

Questo prende lo stesso conto, gia' fatto e gia' verificato, e lo mette in un
foglio di calcolo con le parti separate.

COSA C'E' DENTRO, E COSA NO
Ogni riga e' una VARIANTE, non un prodotto: e' li' che i numeri cambiano (una
cuccia grande costa il doppio della piccola). I prodotti senza costo quotato
ci sono lo stesso, con la nota che dice perche': meglio una riga che dice «non
lo so» di un elenco che sembra completo e non lo e'.

TRE FOGLI
  Listino     una riga per variante, con le parti del costo separate
  Per tipo    la stessa cosa raggruppata, per vedere dove sta il problema
  Come si legge  le ipotesi. Senza, i numeri non si possono usare per
              decidere: l'IVA Printful e' LETTA da un ordine vero, quella
              Printify e' una RISERVA scelta dalla proprietaria, e sono due
              cose molto diverse.

USO
    python3 scripts/perla-verifica-margini.py --json out-listino/margini.json
    python3 scripts/perla-foglio-margini.py
    python3 scripts/perla-foglio-margini.py --json altro.json --dest altro.xlsx
"""
import argparse
import collections
import json
import os
import subprocess
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

QUI = os.path.dirname(os.path.abspath(__file__))
RADICE = os.path.dirname(QUI)
USCITA = os.path.join(RADICE, "out-listino")
NEGOZIO = "https://perlaitaly.com"
SOGLIA = 20.0

# I colori del marchio, gli stessi del tema (config/settings_data.json).
INCHIOSTRO = "FF132A4A"
CREMA = "FFF3E9DA"
ORO = "FFC8862B"
ROSSO = "FFA8442A"
VERDE = "FF2E6A4F"

INTESTAZIONI = [
    ("Linea", 10), ("Tipo", 13), ("Prodotto", 30), ("Variante", 18),
    ("Prezzo", 10), ("Costo prodotto", 15), ("Spedizione", 12),
    ("Imposta", 10), ("Costo totale", 13), ("Guadagno", 11),
    ("Margine", 10), ("Fornitore", 11), ("Nota", 34),
]


def catalogo_pubblico():
    """Tutti i prodotti attivi, per non lasciarne fuori nessuno dal foglio."""
    prodotti = []
    for pagina in (1, 2, 3):
        r = subprocess.run(
            ["curl", "-sfL", "-m", "90",
             "%s/products.json?limit=250&page=%d" % (NEGOZIO, pagina)],
            capture_output=True, text=True)
        if r.returncode != 0:
            break
        try:
            parte = json.loads(r.stdout).get("products", [])
        except ValueError:
            break
        if not parte:
            break
        prodotti += parte
    return prodotti


def linea_di(p):
    grezzi = p.get("tags") or []
    if isinstance(grezzi, str):
        grezzi = grezzi.split(",")
    if any(str(t).strip().endswith("-eu") for t in grezzi):
        return "Europa"
    return "America"


def intesta(ws):
    grassetto = Font(bold=True, color="FFFFFFFF", size=11)
    sfondo = PatternFill("solid", fgColor=INCHIOSTRO)
    for i, (nome, larghezza) in enumerate(INTESTAZIONI, 1):
        c = ws.cell(row=1, column=i, value=nome)
        c.font = grassetto
        c.fill = sfondo
        c.alignment = Alignment(vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = larghezza
    ws.row_dimensions[1].height = 26
    ws.freeze_panes = "A2"


def scrivi_listino(ws, righe):
    intesta(ws)
    filo = Side(style="thin", color="FFE2D8C8")
    bordo = Border(bottom=filo)
    for n, r in enumerate(righe, 2):
        valori = [
            r["linea"], r["tipo"], r["prodotto"], r["taglia"],
            r.get("prezzo"), r.get("costo_prodotto"), r.get("spedizione"),
            r.get("imposta"), r.get("costo"), r.get("guadagno"),
            None if r.get("margine") is None else r["margine"] / 100.0,
            r.get("fornitore"), r.get("nota") or "",
        ]
        for i, v in enumerate(valori, 1):
            c = ws.cell(row=n, column=i, value=v)
            c.border = bordo
            if 5 <= i <= 10:
                c.number_format = '#,##0.00\\ "€"'
            if i == 11:
                c.number_format = "0.0%"
        # Il margine sotto soglia si vede senza cercarlo. Le deroghe no: sono
        # decisioni prese, non allarmi, e colorarle di rosso le farebbe
        # sembrare difetti ogni volta che si riapre il foglio.
        m = r.get("margine")
        if m is not None and m < SOGLIA:
            colore = ORO if r.get("deroga") else ROSSO
            ws.cell(row=n, column=11).font = Font(bold=True, color=colore)
        if r.get("margine") is None:
            ws.cell(row=n, column=13).font = Font(italic=True, color="FF7B8395")
    ws.auto_filter.ref = "A1:%s%d" % (get_column_letter(len(INTESTAZIONI)), len(righe) + 1)


def scrivi_per_tipo(ws, righe):
    quotate = [r for r in righe if r.get("margine") is not None]
    gruppi = collections.OrderedDict()
    for r in sorted(quotate, key=lambda x: (x["linea"], x["tipo"])):
        gruppi.setdefault((r["linea"], r["tipo"]), []).append(r)

    testa = ["Linea", "Tipo", "Varianti", "Prezzo medio", "Costo medio",
             "Spedizione media", "Guadagno medio", "Margine piu' basso",
             "Margine medio"]
    grassetto = Font(bold=True, color="FFFFFFFF", size=11)
    sfondo = PatternFill("solid", fgColor=INCHIOSTRO)
    for i, nome in enumerate(testa, 1):
        c = ws.cell(row=1, column=i, value=nome)
        c.font = grassetto
        c.fill = sfondo
        c.alignment = Alignment(vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = 16 if i > 2 else 13
    ws.row_dimensions[1].height = 26
    ws.freeze_panes = "A2"

    def media(v):
        return sum(v) / float(len(v)) if v else None

    n = 1
    for (linea, tipo), gr in gruppi.items():
        n += 1
        margini = [r["margine"] for r in gr]
        valori = [linea, tipo, len(gr),
                  media([r["prezzo"] for r in gr]),
                  media([r["costo"] for r in gr]),
                  media([r["spedizione"] for r in gr]),
                  media([r["guadagno"] for r in gr]),
                  min(margini) / 100.0, media(margini) / 100.0]
        for i, v in enumerate(valori, 1):
            c = ws.cell(row=n, column=i, value=v)
            if 4 <= i <= 7:
                c.number_format = '#,##0.00\\ "€"'
            if i >= 8:
                c.number_format = "0.0%"
        if min(margini) < SOGLIA:
            ws.cell(row=n, column=8).font = Font(bold=True, color=ROSSO)


def scrivi_note(ws, dati, righe):
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 88
    titolo = Font(bold=True, size=12, color=INCHIOSTRO)
    quotate = [r for r in righe if r.get("margine") is not None]
    senza = [r for r in righe if r.get("margine") is None]
    voci = [
        ("Che cosa e' una riga",
         "Una VARIANTE, non un prodotto: e' li' che i numeri cambiano. La stessa "
         "cuccia in due taglie ha due costi molto diversi."),
        ("Il prezzo",
         "Quello che paga il cliente sul sito, IVA compresa. La spedizione al "
         "cliente e' sempre gratuita: e' quello che promette l'informativa, ed e' "
         "il motivo per cui la spedizione sta nel COSTO e non nel prezzo."),
        ("Il costo del prodotto",
         "Quello che chiede lo stampatore per fabbricarlo. Letto dalle API dei "
         "fornitori nel momento in cui il foglio e' stato costruito, non "
         "trascritto a mano."),
        ("La spedizione",
         "Quella che il negozio paga allo stampatore per mandare UN pezzo. Verso "
         "l'Italia per i prodotti europei, verso gli Stati Uniti per quelli "
         "americani: sono i due casi che succedono davvero, perche' il tema "
         "mostra la linea europea solo a chi guarda dall'Europa."),
        ("L'imposta, e qui le due linee sono diverse",
         "Sui prodotti EUROPEI (Printful) e' l'IVA italiana al 22%, ed e' un "
         "numero LETTO: sta scritto sul preventivo del fornitore. Il negozio non "
         "ha partita IVA, quindi la paga e non la recupera: e' costo a tutti gli "
         "effetti.\n"
         "Sui prodotti AMERICANI (Printify) e' una RISERVA del 5% scelta dalla "
         "proprietaria, non un importo letto: quella merce e' stampata negli "
         "Stati Uniti e consegnata a un indirizzo americano, in Europa non entra "
         "mai, quindi non c'e' IVA italiana. Quello che puo' arrivare e' la sales "
         "tax americana, che dipende dallo stato. La chiudera' la prima fattura "
         "Printify vera."),
        ("Il guadagno",
         "Prezzo meno costo totale. E' quello che resta prima delle commissioni "
         "di pagamento (Shopify Payments trattiene circa l'1,5% + 0,25 EUR a "
         "transazione, e NON e' contato qui) e prima di qualunque spesa fissa."),
        ("Il margine",
         "Guadagno diviso prezzo. La regola decisa e' il 20%: sotto quella soglia "
         "la cella e' rossa. In oro invece ci sono le deroghe, cioe' i prezzi "
         "decisi apposta sotto soglia -- non sono errori, e la colonna Nota dice "
         "quale decisione e quando."),
        ("Le righe senza numeri",
         "%d varianti non hanno un costo quotato: il loro tipo non e' fra quelli "
         "che i due fornitori sanno preventivare da qui. Sono nel foglio lo "
         "stesso, con la nota che lo dice." % len(senza)),
        ("Il cambio",
         "I costi Printify arrivano in dollari e sono convertiti a 1 USD = %.4f "
         "EUR, letto il %s." % (dati.get("cambio", 0), dati.get("aggiornato", "?"))),
        ("Quando e' stato fatto",
         "Il foglio fotografa un momento. I costi dei fornitori cambiano, il "
         "cambio pure: si rifa' con due comandi, e sono scritti in cima a "
         "scripts/perla-foglio-margini.py."),
    ]
    n = 0
    for chiave, testo in voci:
        n += 1
        a = ws.cell(row=n, column=1, value=chiave)
        a.font = titolo
        a.alignment = Alignment(vertical="top", wrap_text=True)
        b = ws.cell(row=n, column=2, value=testo)
        b.alignment = Alignment(vertical="top", wrap_text=True)
        ws.row_dimensions[n].height = max(30, 14 * (testo.count("\n") + 1 + len(testo) // 95))
    n += 2
    ws.cell(row=n, column=1, value="Varianti nel foglio").font = titolo
    ws.cell(row=n, column=2,
            value="%d in tutto: %d con i numeri, %d senza costo quotato."
                  % (len(righe), len(quotate), len(senza)))


def main():
    a = argparse.ArgumentParser()
    a.add_argument("--json", default=os.path.join(USCITA, "margini.json"),
                   help="uscita di perla-verifica-margini.py --json")
    a.add_argument("--dest", default=os.path.join(USCITA, "perla-listino-margini.xlsx"))
    opz = a.parse_args()

    if not os.path.exists(opz.json):
        sys.exit("Manca %s. Prima:\n"
                 "  python3 scripts/perla-verifica-margini.py --json %s"
                 % (opz.json, opz.json))
    dati = json.load(open(opz.json, encoding="utf-8"))
    quotate = {(r["prodotto"], r["taglia"]): r for r in dati["righe"]}

    righe = []
    for p in catalogo_pubblico():
        linea = linea_di(p)
        for v in p["variants"]:
            r = quotate.get((p["title"], v["title"]))
            if r:
                nota = ""
                if r.get("deroga"):
                    nota = "Prezzo sotto soglia deciso apposta: %s" % r["deroga"]
                righe.append(dict(r, linea=linea, nota=nota))
            else:
                righe.append({
                    "linea": linea, "tipo": p.get("product_type", ""),
                    "prodotto": p["title"], "taglia": v["title"],
                    "prezzo": float(v["price"]), "costo_prodotto": None,
                    "spedizione": None, "imposta": None, "costo": None,
                    "guadagno": None, "margine": None,
                    "fornitore": "", "deroga": None,
                    "nota": "Costo non quotato dai fornitori: margine da calcolare a mano.",
                })

    righe.sort(key=lambda r: (r["margine"] is None, r.get("margine") or 0,
                              r["prodotto"], r["taglia"]))

    wb = Workbook()
    scrivi_listino(wb.active, righe)
    wb.active.title = "Listino"
    scrivi_per_tipo(wb.create_sheet("Per tipo"), righe)
    scrivi_note(wb.create_sheet("Come si legge"), dati, righe)

    os.makedirs(os.path.dirname(opz.dest) or ".", exist_ok=True)
    wb.save(opz.dest)
    quotate_n = sum(1 for r in righe if r.get("margine") is not None)
    print("%s\n%d varianti: %d con i numeri, %d senza costo quotato."
          % (opz.dest, len(righe), quotate_n, len(righe) - quotate_n))
    sotto = [r for r in righe if r.get("margine") is not None
             and r["margine"] < SOGLIA and not r.get("deroga")]
    print("Sotto il %.0f%% senza deroga: %d" % (SOGLIA, len(sotto)))


if __name__ == "__main__":
    main()
